import asyncio
import contextlib
import hashlib
import json
import logging
import math
import re
import shutil
import tempfile
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from app.errors import IngestionError, ingestion_error_message
from app.models import Metadata, Segment, VideoObservation, segment_id

log = logging.getLogger(__name__)
DOCUMENTS = {".txt", ".md", ".pdf", ".docx", ".html", ".htm", ".csv", ".json", ".srt", ".vtt"}
VIDEO = {".mp4", ".mov", ".mkv", ".webm"}
AUDIO = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac"}
SUPPORTED = DOCUMENTS | VIDEO | AUDIO


def read_document(path: Path) -> list[tuple[int, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(PdfReader(path).pages)]
    elif suffix == ".docx":
        document = Document(path)
        text = "\n\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(
            " | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows
        )
        pages = [(1, text)]
    else:
        text = path.read_text(encoding="utf-8-sig")
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(text, "html.parser")
            for item in soup(["script", "style"]):
                item.decompose()
            text = soup.get_text("\n")
        pages = [(1, text)]
    if sum(len(t) for _, t in pages) > 2_000_000:
        raise ValueError("Документ превышает лимит 2 млн символов; разделите его")
    units = []
    for page, text in pages:
        for paragraph in re.split(r"\n\s*\n|(?<=[.!?])\s+(?=[А-ЯA-Z])", text):
            paragraph = paragraph.strip()
            for start in range(0, len(paragraph), 1600):
                piece = paragraph[start : start + 1600].strip()
                if piece:
                    units.append((page, piece))
    if not units:
        raise ValueError("Текст не найден. Для сканированного PDF сначала выполните OCR")
    return units


def semantic_merge(units, vectors, threshold=0.65, max_chars=2600):
    chunks = []
    for i, (page, text) in enumerate(units):
        similarity = 0.0
        if i:
            a, b = vectors[i - 1], vectors[i]
            norm = math.sqrt(sum(x * x for x in a) * sum(x * x for x in b))
            similarity = sum(x * y for x, y in zip(a, b, strict=True)) / norm if norm else 0
        if (
            chunks
            and chunks[-1][0] == page
            and similarity >= threshold
            and len(chunks[-1][1]) + len(text) + 2 <= max_chars
        ):
            chunks[-1] = (page, chunks[-1][1] + "\n\n" + text)
        else:
            chunks.append((page, text))
    return chunks


async def command(*args, timeout=180):
    try:
        process = await asyncio.create_subprocess_exec(
            *map(str, args), stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE
        )
    except FileNotFoundError as exc:
        raise IngestionError(
            "FFmpeg or FFprobe was not found. Install both tools, set FFMPEG and FFPROBE in .env, then restart the server."
        ) from exc
    try:
        stdout, stderr = await asyncio.wait_for(process.communicate(), timeout)
    except BaseException:
        if process.returncode is None:
            process.kill()
        await process.wait()
        raise
    if process.returncode:
        # Don't expose file paths / full FFmpeg diagnostic output through the API.
        raise IngestionError(
            "FFmpeg could not read this media file. Check that the file is complete and its format is supported."
        )
    return stdout


async def extract_frames(settings, path, root, index, start, end):
    """Decode once per window instead of starting a decoder for each frame.

    Times are the nominal resampled FPS grid, not measured puck-contact times.
    Include the end of the window explicitly so the last fraction isn't dropped.
    """
    prefix = f"frame-{index:05d}"
    await command(
        settings.ffmpeg,
        "-v",
        "error",
        "-y",
        "-ss",
        start,
        "-i",
        path,
        "-t",
        end - start,
        "-vf",
        f"fps={settings.video_fps}:start_time=0,scale=1280:-2",
        "-q:v",
        "3",
        root / f"{prefix}-%04d.jpg",
    )
    frames = [
        (round(start + i / settings.video_fps, 3), p)
        for i, p in enumerate(sorted(root.glob(f"{prefix}-*.jpg")))
        if start + i / settings.video_fps < end
    ]
    last_time = round(max(start, end - 0.15), 3)
    if not frames or last_time - frames[-1][0] > 0.2:
        target = root / f"{prefix}-tail.jpg"
        await command(
            settings.ffmpeg,
            "-v",
            "error",
            "-y",
            "-ss",
            last_time,
            "-i",
            path,
            "-frames:v",
            "1",
            "-vf",
            "scale=1280:-2",
            target,
        )
        if target.exists():
            frames.append((last_time, target))
    if not frames:
        raise IngestionError("No frames could be decoded from this video window.")
    return frames


class Ingestor:
    def __init__(self, settings, ai, index, storage):
        self.settings, self.ai, self.index, self.storage = settings, ai, index, storage
        self.media_gate = asyncio.Semaphore(3)
        self.queue_lock = asyncio.Lock()
        self.active_id = None
        self.active_task = None

    async def delete_content(self, content_id):
        """Stop ingestion before removing vectors, files, checkpoints and SQL data."""
        async with self.queue_lock:
            content = await self.storage.content(content_id)
            if not content:
                return False
            uploads = (self.settings.data_dir / "uploads").resolve()
            path = Path(content["path"]).resolve()
            checkpoints = (self.settings.data_dir / "checkpoints").resolve()
            cache = (checkpoints / content_id).resolve()
            # Validate the resolved targets before any filesystem removal.
            if path.parent != uploads or cache.parent != checkpoints:
                raise ValueError("Source path is outside the upload directory")
            await self.storage.mark_deleting(content_id)
            if self.active_id == content_id and self.active_task:
                self.active_task.cancel()
                with contextlib.suppress(asyncio.CancelledError, Exception):
                    await self.active_task
            await self.index.delete(content_id)
            await asyncio.to_thread(path.unlink, missing_ok=True)
            if cache.exists():
                await asyncio.to_thread(shutil.rmtree, cache)
            await self.storage.delete_content(content_id)
            return True

    async def progress(self, content_id, message):
        await self.storage.run("UPDATE contents SET progress=? WHERE id=?", (message, content_id))

    async def process(self, content):
        self.ai.require_key()
        content_id = content["id"]
        path = Path(content["path"])
        metadata = Metadata.model_validate_json(content["metadata"])
        if path.suffix.lower() not in DOCUMENTS:
            missing = [
                name
                for name, binary in [
                    ("FFmpeg", self.settings.ffmpeg),
                    ("FFprobe", self.settings.ffprobe),
                ]
                if not shutil.which(binary)
            ]
            if missing:
                raise IngestionError(
                    f"Missing media tools: {', '.join(missing)}. Set FFMPEG and FFPROBE in .env and restart the server."
                )
        if path.suffix.lower() in DOCUMENTS:
            await self.progress(content_id, "Extracting text and semantic boundaries")
            units = await asyncio.to_thread(read_document, path)
            vectors = await self.ai.embed([t for _, t in units])
            chunks = semantic_merge(units, vectors)
            segments = [
                Segment(
                    **metadata.model_dump(),
                    content_id=content_id,
                    segment_id=segment_id(content_id, i),
                    segment_index=i,
                    content_type="document",
                    source=content["name"],
                    transcript=text,
                    page=page,
                    index_version=self.settings.index_version,
                )
                for i, (page, text) in enumerate(chunks)
            ]
        else:
            segments = await self.media(content, metadata)
        await self.progress(content_id, "Embedding analyzed segments")
        all_vectors = await self.ai.embed([s.text() for s in segments])
        await self.index.delete(content_id)
        for offset in range(0, len(segments), 64):
            await self.progress(content_id, f"Indexing {offset}/{len(segments)} segments")
            batch = segments[offset : offset + 64]
            await self.index.upsert(batch, all_vectors[offset : offset + len(batch)])
        await self.index.publish(content_id)
        await self.storage.run(
            "UPDATE contents SET status='ready',progress='Ready',segments=?,error=NULL WHERE id=?",
            (len(segments), content_id),
        )

    async def media(self, content, metadata):
        path = Path(content["path"])
        probe = json.loads(
            await command(
                self.settings.ffprobe,
                "-v",
                "error",
                "-show_format",
                "-show_streams",
                "-of",
                "json",
                path,
            )
        )
        duration = float(probe["format"]["duration"])
        if not math.isfinite(duration) or not 0 < duration <= self.settings.max_media_seconds:
            raise ValueError("Длительность файла превышает лимит или не определена")
        has_audio = any(s["codec_type"] == "audio" for s in probe["streams"])
        is_video = any(s["codec_type"] == "video" for s in probe["streams"])
        if not has_audio and not is_video:
            raise ValueError("В файле нет аудио или видео")
        windows = [
            (float(t), min(t + self.settings.segment_seconds, duration))
            for t in range(0, math.ceil(duration), self.settings.segment_seconds)
            if t < duration
        ]
        await self.progress(content["id"], f"Transcribing and analyzing {len(windows)} segments")
        fingerprint = hashlib.sha256(
            json.dumps(
                {
                    "pipeline": "dense-video-v2",
                    "size": path.stat().st_size,
                    "mtime": path.stat().st_mtime_ns,
                    "model": self.settings.vision_model,
                    "fps": self.settings.video_fps,
                    "seconds": self.settings.segment_seconds,
                    "overlap": self.settings.frame_overlap_seconds,
                    "metadata": metadata.model_dump(),
                    "audio_events": self.settings.audio_events,
                    "audio_model": self.settings.audio_model,
                },
                sort_keys=True,
            ).encode()
        ).hexdigest()[:20]
        cache = self.settings.data_dir / "checkpoints" / content["id"] / fingerprint
        cache.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="touchline-") as temp:
            root = Path(temp)

            async def extract_audio(start, end, target):
                async with self.media_gate:
                    await command(
                        self.settings.ffmpeg,
                        "-v",
                        "error",
                        "-y",
                        "-ss",
                        start,
                        "-i",
                        path,
                        "-t",
                        end - start,
                        "-vn",
                        "-ac",
                        "1",
                        "-ar",
                        "16000",
                        target,
                    )

            async def asr():
                if not has_audio:
                    return []
                results = []
                # Ten minutes of mono 16kHz PCM stays below Whisper's 25MB input limit.
                for offset in range(0, math.ceil(duration), 600):
                    cache_path = cache / f"speech-{offset}.json"
                    if cache_path.exists():
                        results.extend(json.loads(cache_path.read_text(encoding="utf-8")))
                        continue
                    target = root / f"asr-{offset}.wav"
                    await extract_audio(offset, min(offset + 600, duration), target)
                    block = await self.ai.transcribe(target, offset)
                    cache_path.write_text(json.dumps(block, ensure_ascii=False), encoding="utf-8")
                    results.extend(block)
                    target.unlink(missing_ok=True)
                return results

            async def describe(i, start, end):
                cache_path = cache / f"vision-{i}.json"
                analysis_start = max(0, start - self.settings.frame_overlap_seconds)
                analysis_end = min(duration, end + self.settings.frame_overlap_seconds)
                frame_times = []
                observation = None
                if is_video and cache_path.exists():
                    cached = json.loads(cache_path.read_text(encoding="utf-8"))
                    observation = VideoObservation.model_validate(cached["observation"])
                    frame_times = cached["frame_times"]
                elif is_video:
                    async with self.media_gate:
                        frames = await extract_frames(
                            self.settings, path, root, i, analysis_start, analysis_end
                        )
                    frame_times = [t for t, _ in frames]
                    observation = await self.ai.vision(frames, metadata.model_dump())
                    cache_path.write_text(
                        json.dumps(
                            {"observation": observation.model_dump(), "frame_times": frame_times},
                            ensure_ascii=False,
                        ),
                        encoding="utf-8",
                    )
                    for _, target in frames:
                        target.unlink(missing_ok=True)
                audio = ""
                if self.settings.audio_events and has_audio:
                    target = root / f"events-{i}.wav"
                    await extract_audio(start, end, target)
                    audio = await self.ai.audio_events(target)
                    target.unlink(missing_ok=True)
                return observation, audio, frame_times

            async def descriptions():
                results = []
                # Bound queued tasks, decoded frames and base64 memory, not just network calls.
                for offset in range(0, len(windows), self.settings.model_concurrency):
                    batch = windows[offset : offset + self.settings.model_concurrency]
                    async with asyncio.TaskGroup() as group:
                        tasks = [
                            group.create_task(describe(offset + i, a, b))
                            for i, (a, b) in enumerate(batch)
                        ]
                    results.extend(task.result() for task in tasks)
                    await self.progress(
                        content["id"], f"Analyzing segments {len(results)}/{len(windows)}"
                    )
                return results

            # TaskGroup cancels sibling tasks on failure before temporary files are removed.
            async with asyncio.TaskGroup() as group:
                speech_task = group.create_task(asr())
                descriptions_task = group.create_task(descriptions())
            speech, descriptions_result = speech_task.result(), descriptions_task.result()
        segments = []
        for i, ((start, end), (observation, audio, frame_times)) in enumerate(
            zip(windows, descriptions_result, strict=True)
        ):
            transcript = " ".join(
                s["text"] for s in speech if s["start"] < end and s["end"] > start
            )
            segments.append(
                Segment(
                    **metadata.model_dump(),
                    content_id=content["id"],
                    segment_id=segment_id(content["id"], i),
                    segment_index=i,
                    source=content["name"],
                    content_type="video" if is_video else "audio",
                    start_time=start,
                    end_time=end,
                    transcript=transcript,
                    visual_description=observation.description if observation else "",
                    sport=observation.sport if observation else None,
                    scores=observation.scores if observation else [],
                    events=observation.events if observation else [],
                    frame_times=frame_times,
                    ingestion_version=f"dense-video-v2:{self.settings.segment_seconds}s:{self.settings.video_fps}fps",
                    audio_description=audio,
                    index_version=self.settings.index_version,
                )
            )
        return segments

    async def worker(self):
        # A crash during deletion must never put the source back in the queue.
        pending = await self.storage.run(
            "SELECT id FROM contents WHERE status='deleting'", fetch=True
        )
        for row in pending:
            try:
                await self.delete_content(row["id"])
            except Exception as exc:
                log.error("deletion_recovery_failed type=%s", type(exc).__name__)
        while True:
            async with self.queue_lock:
                rows = await self.storage.run(
                    "SELECT * FROM contents WHERE status='queued' ORDER BY created_at LIMIT 1",
                    fetch=True,
                )
                if rows:
                    content = rows[0]
                    await self.storage.run(
                        "UPDATE contents SET status='running',error=NULL WHERE id=?",
                        (content["id"],),
                    )
                    self.active_id = content["id"]
                    task = self.active_task = asyncio.create_task(self.process(content))
            if not rows:
                await asyncio.sleep(1)
                continue
            try:
                async with asyncio.timeout(6 * 3600):
                    await task
            except asyncio.CancelledError:
                if asyncio.current_task().cancelling():
                    raise
            except Exception as exc:
                # No input, secrets or upstream response bodies in logs.
                log.error(
                    "ingestion_failed content_id=%s type=%s", content["id"], type(exc).__name__
                )
                await self.storage.run(
                    "UPDATE contents SET status='failed',progress='Processing failed',error=? WHERE id=? AND status='running'",
                    (
                        ingestion_error_message(exc),
                        content["id"],
                    ),
                )
            finally:
                self.active_id = self.active_task = None
